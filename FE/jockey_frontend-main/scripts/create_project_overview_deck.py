from pathlib import Path
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "docs" / "Philippine-Jockey-Product-Flows-with-Mobile-Screenshots.pptx"
OUT.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

GREEN = RGBColor(15, 61, 48)
GREEN2 = RGBColor(28, 86, 69)
GOLD = RGBColor(198, 159, 83)
CREAM = RGBColor(246, 242, 232)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(29, 43, 39)
MUTED = RGBColor(105, 119, 112)
PALE = RGBColor(231, 237, 231)
LINE = RGBColor(220, 224, 216)
ORANGE = RGBColor(178, 117, 56)


def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
            margin=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=True, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def base(title, eyebrow="PHILIPPINE JOCKEY  /  PRODUCT OVERVIEW", note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = CREAM
    textbox(slide, .55, .24, 8, .25, eyebrow, 9, ORANGE, True)
    textbox(slide, .55, .58, 12.2, .6, title, 27, GREEN, True)
    rect(slide, .55, 1.27, 12.2, .025, GOLD, False)
    textbox(slide, .55, 7.12, 10, .2,
            note or "Play-money demo · Mobile app and web · Philippine Peso (₱)",
            8, MUTED)
    textbox(slide, 12.15, 7.08, .55, .24, f"{len(prs.slides):02}", 9, MUTED,
            align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, title, body, accent=GREEN, title_size=15, body_size=11):
    rect(slide, x, y, w, h, WHITE, line=LINE)
    rect(slide, x, y, .08, h, accent, False)
    textbox(slide, x + .18, y + .12, w - .34, .38, title, title_size, accent, True)
    textbox(slide, x + .18, y + .53, w - .34, h - .62, body, body_size, INK,
            valign=MSO_ANCHOR.TOP)


def flow(slide, items, y=3.05, x=.75, total_w=11.85, h=1.02, gap=.32,
         font_size=13):
    n = len(items)
    cw = (total_w - gap * (n - 1)) / n
    for i, label in enumerate(items):
        xx = x + i * (cw + gap)
        rect(slide, xx, y, cw, h, GREEN if i in (0, n - 1) else WHITE,
             line=None if i in (0, n - 1) else LINE)
        textbox(slide, xx + .08, y + .08, cw - .16, h - .16, label,
                font_size, WHITE if i in (0, n - 1) else GREEN, True,
                align=PP_ALIGN.CENTER)
        if i < n - 1:
            textbox(slide, xx + cw, y + .21, gap, .5, "›", 26, GOLD, True,
                    align=PP_ALIGN.CENTER)


# 1 — Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = GREEN
rect(slide, 0, 0, .2, 7.5, GOLD, False)
textbox(slide, .8, .75, 5, .35, "PHILIPPINE JOCKEY", 12, GOLD, True)
textbox(slide, .8, 1.55, 8.7, 1.5, "Product flows\n& page guide", 36, WHITE, True)
textbox(slide, .85, 3.35, 7.8, .75,
        "A walkthrough of the mobile app and web experience—from sign-in to race day, bets, results, and account tools.",
        18, RGBColor(220, 232, 224))
rect(slide, .85, 5.18, 3.4, .55, GOLD)
textbox(slide, 1.02, 5.22, 3.1, .44, "APP + WEB · DEMO EDITION", 11, GREEN, True,
        align=PP_ALIGN.CENTER)
textbox(slide, .85, 6.74, 9, .3, "Product overview · October 2026 · Currency: PHP (₱)", 10,
        RGBColor(220, 232, 224))

# 2 — Purpose and audiences
slide = base("What the product does")
card(slide, .7, 1.65, 3.85, 1.75, "Discover racing", "Browse race days, clubs, race cards, horses, jockeys, tips, and live race information.", GREEN)
card(slide, 4.75, 1.65, 3.85, 1.75, "Place demo bets", "Choose a race and game, select runners, set a stake, review the ticket, and see confirmation.", GOLD)
card(slide, 8.8, 1.65, 3.85, 1.75, "Follow outcomes", "Check results, podiums, dividends, tickets, win notifications, favorites, and profile tools.", GREEN2)
textbox(slide, .8, 3.82, 11.5, .4, "Designed as one racing experience across", 15, MUTED)
flow(slide, ["Expo mobile app", "Shared domain UI/data", "Next.js web app"], y=4.35, h=.85)
textbox(slide, .8, 5.65, 11.4, .65,
        "Important: this repository is a frontend demo. It currently uses mock data and simulated wallet/betting behavior; no real money is accepted or paid out.",
        14, ORANGE, True, valign=MSO_ANCHOR.TOP)

# 3 — Entry and navigation map
slide = base("How people move through the app")
flow(slide, ["Splash", "Onboarding", "Welcome / Sign in / Sign up", "Home dashboard"], y=1.78, h=.92, font_size=12)
textbox(slide, .8, 3.05, 5.7, .35, "PRIMARY APP TABS", 10, ORANGE, True)
for i, (title, desc) in enumerate([
    ("Home", "Dashboard, live feature, today's races, shortcuts, My Bets"),
    ("Races", "Meeting calendar, clubs, filters, race cards"),
    ("Live", "Stream preview, updates, entries, totalizer, available games"),
    ("Results", "Official/provisional/cancelled result list and details"),
    ("Profile", "Demo balance, edit profile, history, favorites, notifications, settings"),
]):
    yy = 3.48 + i * .57
    rect(slide, .8, yy, .95, .4, GREEN)
    textbox(slide, .84, yy + .01, .87, .36, title, 10, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(slide, 1.95, yy, 4.8, .4, desc, 11, INK)
card(slide, 7.05, 3.05, 5.35, 2.9, "Shared secondary pages", "Race detail and entries\nBet selection and cart\nReview and confirmation\nTicket detail and betting history\nHorse / jockey / club profiles\nGames, guides, rewards, totalizer", accent=GOLD, body_size=13)

# 4 — Account journey
slide = base("Flow 1 · Account access")
flow(slide, ["Splash", "Onboarding", "Welcome", "Log in or sign up", "Home"], y=1.8, h=1.0, font_size=12)
card(slide, .8, 3.35, 5.7, 2.2, "Sign up", "Collects first name, last name, email, mobile, and password. The demo creates a client-side session and sample account.", accent=GREEN)
card(slide, 6.8, 3.35, 5.7, 2.2, "Sign in / sign out", "Sign in opens the app session. Profile offers log out. Auth routes are shared by mobile and web; current demo mode can bypass authentication.", accent=GOLD)
textbox(slide, .85, 5.95, 11.4, .45, "Routes: /splash · /(onboarding)/onboarding · /(auth)/welcome · /(auth)/login · /(auth)/signup", 11, MUTED)

# 5 — Dashboard and race discovery
slide = base("Flow 2 · Dashboard to race card")
flow(slide, ["Dashboard", "Race schedule", "Choose date / club", "Race detail", "Entries"], y=1.85, h=.98, font_size=12)
card(slide, .8, 3.25, 3.8, 2.2, "Dashboard entry points", "Live race feature\nToday's race cards\nPromoted games\nMy Bets\nShortcuts to guides and utilities", accent=GREEN, body_size=12)
card(slide, 4.8, 3.25, 3.8, 2.2, "Race discovery", "Pick a meeting date and racing club. Filter by race status or search. Tap a card to inspect the race.", accent=GOLD, body_size=12)
card(slide, 8.8, 3.25, 3.8, 2.2, "Race detail", "See venue, race conditions, entries, participants, game markets, tips, and routes into betting.", accent=GREEN2, body_size=12)
textbox(slide, .85, 5.85, 11.4, .45, "Routes: /(tabs)/index · /(tabs)/races · /race/[id] · /race/[id]/entries · /horse/[id] · /jockey/[id] · /club/[slug]", 10, MUTED)

# 6 — Betting journey
slide = base("Flow 3 · Place a demo bet")
flow(slide, ["Race", "Market", "Runner selection", "Bet cart", "Review", "Confirmation", "Ticket"], y=1.83, h=1.0, font_size=11)
card(slide, .8, 3.3, 3.8, 2.25, "1 · Select", "Choose a race and supported market (WIN, PLACE, FORECAST, TRIFECTA, and others). Select runner(s) and enter a stake.", accent=GREEN, body_size=11)
card(slide, 4.8, 3.3, 3.8, 2.25, "2 · Review", "Review the selection, stake, and ticket summary before submitting. This is a simulated play-money action.", accent=GOLD, body_size=11)
card(slide, 8.8, 3.3, 3.8, 2.25, "3 · Track", "The app shows a confirmation and ticket details; tickets can be revisited from My Bets / Betting History.", accent=GREEN2, body_size=11)
textbox(slide, .85, 5.9, 11.4, .4, "Routes: /selection/[id] · /bet-cart · /review · /bet-complete · /ticket/[id] · /betting-history", 11, MUTED)

# 7 — Live race and pools
slide = base("Flow 4 · Live race experience")
flow(slide, ["Live tab", "Stream preview", "Live updates", "Entries", "Totalizer / games"], y=1.9, h=.95, font_size=12)
card(slide, .8, 3.35, 5.65, 2.3, "Live center", "The live tab shows the featured race, a stream preview, update link, race entries, totalizer, and available games. The included stream is demo content.", accent=GREEN, body_size=12)
card(slide, 6.75, 3.35, 5.65, 2.3, "Pools and markets", "Totalizer pages show sample pool/dividend information. Games routes explain available bet types and can start a selection flow.", accent=GOLD, body_size=12)
textbox(slide, .85, 5.92, 11.4, .4, "Routes: /(tabs)/live · /live/stream · /live/updates · /totalizer/[id] · /totalizer/pool/[id] · /totalizer/dividend/[id] · /games · /games/[id]", 10, MUTED)

# 8 — Results and notifications
slide = base("Flow 5 · Race result to winner notification")
flow(slide, ["Results list", "Result detail", "Match ticket to result", "Win notice", "Profile notifications"], y=1.88, h=1.0, font_size=12)
card(slide, .8, 3.35, 5.65, 2.25, "Result detail", "Users can inspect official placements, podium, winning time, dividends, race summary, and horse/jockey profiles.", accent=GREEN, body_size=12)
card(slide, 6.75, 3.35, 5.65, 2.25, "Bet win notification", "When official results are compared with the user's saved tickets, the demo can create a notification describing a winning ticket and its simulated return.", accent=GOLD, body_size=12)
textbox(slide, .85, 5.95, 11.4, .35, "Routes: /(tabs)/results · /result/[id] · /notifications · /ticket/[id]", 11, MUTED)

# 9 — Account, favorites, settings
slide = base("Flow 6 · Profile and account tools")
flow(slide, ["Profile", "Choose account tool", "Edit / review", "Return to profile"], y=1.86, h=.98, font_size=12)
items = [
    ("Balance & withdrawal", "Displays a demo balance. Withdrawal is not processed."),
    ("Betting history", "View tickets and their current demo status."),
    ("Favorites", "Save and revisit horses or contenders."),
    ("Notifications", "Review updates and simulated bet-win notices."),
    ("Settings & profile", "Edit basic profile details and account preferences."),
]
for i, (t, b) in enumerate(items):
    col, row = i % 3, i // 3
    card(slide, .8 + col * 4.05, 3.25 + row * 1.5, 3.75, 1.25, t, b,
         accent=GOLD if i % 2 else GREEN, title_size=13, body_size=10)
textbox(slide, .85, 6.45, 11.4, .3, "Routes: /(tabs)/profile · /withdraw · /edit-profile · /settings · /favorites · /betting-history · /notifications · /rewards", 10, MUTED)

# 10 — Mobile page catalog
slide = base("Mobile app · Page inventory")
groups = [
    ("Entry & auth", "Splash · Onboarding · Welcome · Login · Sign up", GREEN),
    ("Main tabs", "Home · Races · Live · Results · Profile", GOLD),
    ("Racing", "Race detail · Entries · Result detail · Horse profile/statistics · Jockey profile/statistics · Club profile", GREEN2),
    ("Betting", "Selection · Bet cart · Review · Bet complete · Ticket detail · Betting history", GREEN),
    ("Live & info", "Live updates · Stream · Games · Game detail · Racing guide · JockeyPH guide · Contender", GOLD),
    ("Account & pools", "Notifications · Settings · Edit profile · Rewards · Withdraw · Totalizer · Pool · Dividend", GREEN2),
]
for i, (t, b, c) in enumerate(groups):
    col, row = i % 2, i // 2
    card(slide, .8 + col * 6.05, 1.62 + row * 1.65, 5.7, 1.4, t, b, c, title_size=14, body_size=11)

# 11 — Web page catalog
slide = base("Web app · Page inventory")
card(slide, .8, 1.7, 5.7, 1.48, "Public access", "Login · Register", GREEN, body_size=13)
card(slide, 6.8, 1.7, 5.7, 1.48, "Dashboard and racing", "Dashboard · Races · Race detail · Entries · Results · Result detail", GOLD, body_size=12)
card(slide, .8, 3.45, 5.7, 1.48, "Betting and live", "Selection · Bet cart · Review · Bet complete · Betting history · Live · Stream · Updates", GREEN2, body_size=12)
card(slide, 6.8, 3.45, 5.7, 1.48, "Profile and info", "Profile · Edit profile · Settings · Notifications · Favorites · Rewards · Withdraw · Games · Guides · Totalizer · Horse/jockey/club pages", GREEN, body_size=11)
textbox(slide, .9, 5.6, 11.2, .75, "Web dashboard navigation includes Home, Races, Race Program & Tips, JockeyPH Guide, Live, Results, and Profile. A mobile tab bar supports small screens.", 13, INK, valign=MSO_ANCHOR.TOP)

# 12 — Shared data and current implementation status
slide = base("Data, wallet, and integration status")
card(slide, .8, 1.68, 3.8, 2.05, "Current demo seam", "Screens call a shared API client. Mock mode provides local sample races, results, tickets, and profile data.", GREEN, body_size=12)
card(slide, 4.8, 1.68, 3.8, 2.05, "Demo wallet", "Balance is sample UI state. Bet placement, winnings, and withdrawals do not move real funds or connect to a live ledger.", GOLD, body_size=12)
card(slide, 8.8, 1.68, 3.8, 2.05, "Future Django API", "A contract documents expected endpoints, but the backend must implement, secure, and connect persistence before data is database-backed.", GREEN2, body_size=12)
flow(slide, ["Frontend mock data", "API contract", "Django endpoints", "Database + ledger"], y=4.45, h=.9, font_size=12)
textbox(slide, .9, 5.72, 11.2, .7, "A real wallet requires server-side balance and immutable transaction records, authenticated APIs, idempotent operations, and reconciliation. Do not trust client-side balance changes.", 13, ORANGE, True, valign=MSO_ANCHOR.TOP)

# 13 — End-to-end summary
slide = base("The complete customer journey")
flow(slide, ["Sign in", "Discover a race", "Study runners", "Place a demo bet", "Follow live race", "Check official result", "Review ticket & notice"], y=1.85, h=1.08, font_size=10)
card(slide, .9, 3.55, 5.45, 1.95, "User outcome", "One connected journey links discovery, runner information, a bet ticket, race coverage, result detail, and follow-up notification.", GREEN, body_size=13)
card(slide, 6.95, 3.55, 5.45, 1.95, "Product boundary", "The current implementation demonstrates the experience with mock data and play-money values. Database persistence and real payment processing are future backend work.", GOLD, body_size=13)
textbox(slide, .95, 6.05, 11.2, .45, "Thank you · Philippine Jockey · App + Web", 16, GREEN, True, align=PP_ALIGN.CENTER)

# Appendix: phone-sized screenshots for every mobile app route.
manifest_path = OUT.parent / "mobile-screenshots" / "manifest.json"
if manifest_path.exists():
    screens = json.loads(manifest_path.read_text(encoding="utf-8"))
    pages = [screen for screen in screens if (manifest_path.parent / screen["image"]).exists()]
    for offset in range(0, len(pages), 3):
        batch = pages[offset:offset + 3]
        slide = base(f"Mobile screen gallery · pages {offset + 1}–{offset + len(batch)} of {len(pages)}",
                     eyebrow="APPENDIX  /  MOBILE SCREENSHOTS")
        for index, screen in enumerate(batch):
            x = 1.00 + index * 4.45
            textbox(slide, x - .08, 1.39, 2.55, .30, screen["title"], 11, GREEN, True,
                    align=PP_ALIGN.CENTER)
            slide.shapes.add_picture(str(manifest_path.parent / screen["image"]),
                                     Inches(x), Inches(1.72), width=Inches(2.38), height=Inches(5.15))
            textbox(slide, x - .16, 6.90, 2.70, .20, screen["route"], 7, MUTED,
                    align=PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
